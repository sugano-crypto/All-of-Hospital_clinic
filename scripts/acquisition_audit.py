"""Inventory every linked spreadsheet/archive on official source pages.
This is an acquisition audit, not a verified nationwide normalizer or production updater.
Preserves raw documents, all workbook sheets, hashes, failures, and monthly retry state.
"""
import argparse, hashlib, io, json, re, urllib.request, urllib.parse, zipfile
from datetime import datetime
from html.parser import HTMLParser
from pathlib import Path
from zoneinfo import ZoneInfo

SOURCES = json.loads(Path(__file__).with_name('sources.json').read_text())
ALLOWED = {'kouseikyoku.mhlw.go.jp', 'www.mhlw.go.jp'}
class OfficialRedirect(urllib.request.HTTPRedirectHandler):
    def redirect_request(self, req, fp, code, msg, headers, newurl):
        check_url(newurl)
        return super().redirect_request(req, fp, code, msg, headers, newurl)
def check_url(url):
    parsed = urllib.parse.urlsplit(url)
    if parsed.scheme != 'https' or parsed.hostname not in ALLOWED:
        raise ValueError('公式HTTPS以外のURLは取得しません')
def fetch(url):
    check_url(url)
    # Inherits proxy/CA settings. Do not bypass environment policy.
    opener=urllib.request.build_opener(OfficialRedirect())
    with opener.open(urllib.request.Request(url,headers={'User-Agent':'NationalMedicalFacilitySearch/0.1 source-audit'}),timeout=30) as response:
        raw=response.read(100*1024*1024+1)
        if len(raw)>100*1024*1024: raise ValueError('取得ファイルが100MBを超えました')
        return raw
class Links(HTMLParser):
    def __init__(self): super().__init__(); self.links=[]; self.current=None; self.text=[]
    def handle_starttag(self,tag,attrs):
        if tag=='a': self.current=dict(attrs).get('href'); self.text=[]
    def handle_data(self,data):
        if self.current is not None: self.text.append(data)
    def handle_endtag(self,tag):
        if tag=='a' and self.current is not None:
            self.links.append((self.current,''.join(self.text).strip())); self.current=None

def discover(raw,url):
    parser=Links(); parser.feed(raw.decode('utf-8-sig'))
    out={}
    for href,title in parser.links:
        full=urllib.parse.urljoin(url,href)
        if re.search(r'\.(xlsx?|zip)(?:\?|$)',full,re.I):
            check_url(full); out[full]={'url':full,'title':title}
    return list(out.values())

def inspect_excel(raw,name):
    """Audit all sheets, including hidden sheets; never claim semantic completeness."""
    if name.lower().endswith('.xlsx'):
        import openpyxl
        book=openpyxl.load_workbook(io.BytesIO(raw),read_only=True,data_only=False)
        try:
            return [{'sheet':s.title,'rows':s.max_row,'columns':s.max_column,'nonempty_rows':sum(any(v is not None for v in row) for row in s.iter_rows(values_only=True))} for s in book.worksheets]
        finally: book.close()
    import xlrd
    book=xlrd.open_workbook(file_contents=raw,on_demand=True)
    try:
        return [{'sheet':s.name,'rows':s.nrows,'columns':s.ncols,'nonempty_rows':sum(any(v!='' for v in s.row_values(i)) for i in range(s.nrows))} for s in book.sheets()]
    finally: book.release_resources()

def inspect_download(raw,url):
    name=urllib.parse.urlsplit(url).path
    if name.lower().endswith('.zip'):
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            if len(archive.infolist())>2000 or sum(f.file_size for f in archive.infolist())>500*1024*1024: raise ValueError('ZIP展開サイズ超過')
            workbooks=[]
            for f in archive.infolist():
                if re.search(r'\.xlsx?$',f.filename,re.I):
                    workbooks.append({'file':f.filename,'sheets':inspect_excel(archive.read(f),f.filename)})
            if not workbooks: raise ValueError('ZIP内にExcelがありません')
            return workbooks
    return [{'file':name.rsplit('/',1)[-1],'sheets':inspect_excel(raw,name)}]

def pending_sources(state,month,retry):
    if retry and state.get('month')!=month: raise ValueError('同月20日の取得状態がありません。全局取得を先に実行してください。')
    return [s for s in SOURCES if not retry or state.get('bureaus',{}).get(s['id'],{}).get('acquisition_status')!='downloaded']

def run(output,mode,now=None,fetcher=fetch):
    now=now or datetime.now(ZoneInfo('Asia/Tokyo')); month=now.strftime('%Y-%m')
    output=Path(output);output.mkdir(parents=True,exist_ok=True)
    state_path=output/'state.json';state=json.loads(state_path.read_text()) if state_path.exists() else {}
    targets=pending_sources(state,month,mode=='retry')
    if state.get('month')!=month: state={'month':month,'bureaus':{}}
    state.update({'tested_at':now.isoformat(),'ready_to_publish':False,'notice':'原本取得・シート読込の監査です。47都道府県の網羅・全届出項目・正式名称対応は別途検証が必要です。'})
    for source in targets:
        report={'name':source['name'],'acquisition_status':'failed','semantic_validation':'not_verified','pages':{},'files':[],'errors':[]}
        directory=output/month/source['id'];directory.mkdir(parents=True,exist_ok=True)
        documents={}
        for kind in ['facilities','filings']:
            url=source[kind]
            try:
                raw=fetcher(url);(directory/f'{kind}.html').write_bytes(raw)
                links=discover(raw,url)
                if not links: raise ValueError('Excel・ZIPリンクが見つかりません')
                report['pages'][kind]={'url':url,'linked_files':len(links)}
                for link in links: documents[link['url']]=link
            except Exception as e: report['errors'].append({'stage':kind,'url':url,'error':str(e)})
        for url,link in documents.items():
            try:
                raw=fetcher(url);digest=hashlib.sha256(raw).hexdigest();ext=Path(urllib.parse.urlsplit(url).path).suffix
                path=directory/f'{digest}{ext}';path.write_bytes(raw)
                report['files'].append({**link,'sha256':digest,'bytes':len(raw),'workbooks':inspect_download(raw,url)})
            except Exception as e: report['errors'].append({'stage':'file','url':url,'error':str(e)})
        if report['files'] and not report['errors']:report['acquisition_status']='downloaded'
        report['tested_at']=now.isoformat();state['bureaus'][source['id']]=report
        state_path.write_text(json.dumps(state,ensure_ascii=False,indent=2))
    return state
if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--output',default='work/acquisition');parser.add_argument('--mode',choices=['all','retry'],default='all');args=parser.parse_args()
    state=run(args.output,args.mode)
    print(json.dumps({'month':state['month'],'ready_to_publish':False,'bureaus':{k:{'acquisition_status':v['acquisition_status'],'files':len(v['files']),'errors':v['errors']} for k,v in state['bureaus'].items()}},ensure_ascii=False,indent=2))
    raise SystemExit(0 if all(b['acquisition_status']=='downloaded' for b in state['bureaus'].values()) else 1)
